#!/usr/bin/env python3
"""
build_slidesgo_presentation.py
Creates the UAV Disturbance Observer presentation using the user's exact Slidesgo template:
"Sewerage System Construction Project Proposal by Slidesgo.pptx"
- Preserves the authentic Slidesgo visual assets, colors, and styling.
- Populates all slides with our project technical content.
- Replaces image placeholders with real simulation figures.
- Removes all unused boilerplate/icon slides, leaving a clean 7-slide deck.
"""

import os
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TEMPLATE_PATH = "/mnt/c/Users/N V D S SARATCHANDRA/Downloads/Sewerage System Construction Project Proposal by Slidesgo.pptx"
OUT_PATH = "/mnt/d/ROS/UAV_Disturbance_Observer_Slidesgo_Presentation.pptx"
IMG_DIR = "/mnt/d/ROS/docs/images"

def replace_shape_image(slide, shape_id_or_name, img_path):
    """Finds an existing picture shape, removes it, and inserts a new fitted picture in its place."""
    target_shape = None
    for s in list(slide.shapes):
        if s.shape_type == 13 or (shape_id_or_name and shape_id_or_name in s.name):
            target_shape = s
            break

    if not target_shape or not os.path.exists(img_path):
        return None

    l, t, w, h = target_shape.left, target_shape.top, target_shape.width, target_shape.height
    sp = target_shape._element
    sp.getparent().remove(sp)

    # Fit image to bounding box preserving aspect ratio
    with Image.open(img_path) as im:
        im_w, im_h = im.size
        im_ratio = im_w / im_h

    box_ratio = w / h
    if im_ratio > box_ratio:
        fit_w = w
        fit_h = int(w / im_ratio)
    else:
        fit_h = h
        fit_w = int(h * im_ratio)

    offset_x = int(l + (w - fit_w) / 2.0)
    offset_y = int(t + (h - fit_h) / 2.0)

    new_pic = slide.shapes.add_picture(img_path, offset_x, offset_y, fit_w, fit_h)
    return new_pic

def set_text(shape, text, font_size=None, bold=None, color=None):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if font_size:
        p.font.size = Pt(font_size)
    if bold is not None:
        p.font.bold = bold
    if color:
        p.font.color.rgb = color

def delete_slide(prs, index):
    rId = prs.slides._sldIdLst[index].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[index]

def build():
    print(f"Loading Slidesgo template from: {TEMPLATE_PATH}")
    prs = Presentation(TEMPLATE_PATH)
    total_initial = len(prs.slides)
    print(f"Loaded template with {total_initial} slides.")

    # Slides to keep (0-indexed in original template):
    # Slide 1 (0): TITLE
    # Slide 6 (5): TITLE_AND_TWO_COLUMNS
    # Slide 12 (11): TITLE_ONLY (Control Architecture)
    # Slide 14 (13): CUSTOM_4 (Gazebo 3D Arena)
    # Slide 18 (17): CUSTOM_22 (Multi-Gust Benchmark)
    # Slide 19 (18): CUSTOM_22_1 (KPI Dashboard / GCS)
    # Slide 20 (19): CUSTOM_3 (Thanks / Q&A)
    KEEP_INDICES = {0, 5, 11, 13, 17, 18, 19}

    slides_list = list(prs.slides)

    # -------------------------------------------------------------
    # 1. EDIT SLIDE 1 (TITLE)
    # -------------------------------------------------------------
    s1 = slides_list[0]
    print("Configuring Slide 1: Title...")
    for s in s1.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "Sewerage" in txt:
                s.text_frame.text = "Nonlinear Disturbance Observer (DOB) for Quadrotors"
                p = s.text_frame.paragraphs[0]
                p.font.size = Pt(26)
                p.font.bold = True
            elif "presentation begins" in txt:
                s.text_frame.text = "Autonomous Aerial Robotics  |  ROS 2 Humble & Gazebo 11\nAuthor: Nallabati Venkata Durga Sai Sarat Chandra"
                p = s.text_frame.paragraphs[0]
                p.font.size = Pt(11)

    # -------------------------------------------------------------
    # 2. EDIT SLIDE 6 (TITLE_AND_TWO_COLUMNS -> PROBLEM STATEMENT)
    # -------------------------------------------------------------
    s2 = slides_list[5]
    print("Configuring Slide 2: Problem Statement...")
    for s in s2.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "Objectives" in txt:
                s.text_frame.text = "Problem Statement: Wind Shocks vs. Classical PID"
                s.text_frame.paragraphs[0].font.size = Pt(20)
                s.text_frame.paragraphs[0].font.bold = True
            elif "Our aim" in txt:
                s.text_frame.text = "Conventional PID Failure Modes"
                s.text_frame.paragraphs[0].font.size = Pt(14)
                s.text_frame.paragraphs[0].font.bold = True
            elif "The goal" in txt:
                s.text_frame.text = "Active Disturbance Rejection (ADRC)"
                s.text_frame.paragraphs[0].font.size = Pt(14)
                s.text_frame.paragraphs[0].font.bold = True
            elif "An aim in a corporate" in txt:
                s.text_frame.clear()
                points = [
                    "• 8.5 m/s crosswind exerts 11.2 N on a 1.5 kg drone (75% vehicle weight).",
                    "• Severe reaction lag: PID cannot act until large tracking errors occur.",
                    "• Integral windup accumulates error over seconds, flinging drone into obstacles.",
                    "• High PID gains amplify sensor noise, causing motor chattering."
                ]
                for i, pt in enumerate(points):
                    p = s.text_frame.add_paragraph() if i > 0 else s.text_frame.paragraphs[0]
                    p.text = pt
                    p.font.size = Pt(10)
                    p.space_after = Pt(4)
            elif "Goals inside a company" in txt:
                s.text_frame.clear()
                points = [
                    "• Rapid observer convergence: Reconstruct fluid drag in 28.5 ms (ω_obs = 35 rad/s).",
                    "• Feedforward decoupling: a_des = a_PID - (1/m) F_hat cancels drag in acceleration space.",
                    "• Nonlinear tilt compensation: Scales thrust by 1/(cos φ cos θ) to prevent altitude drop.",
                    "• Safety guarantee: Verified clearance ≥ 0.82 m past all facility structures."
                ]
                for i, pt in enumerate(points):
                    p = s.text_frame.add_paragraph() if i > 0 else s.text_frame.paragraphs[0]
                    p.text = pt
                    p.font.size = Pt(10)
                    p.space_after = Pt(4)

    # -------------------------------------------------------------
    # 3. EDIT SLIDE 12 (TITLE_ONLY -> CONTROL ARCHITECTURE)
    # -------------------------------------------------------------
    s3 = slides_list[11]
    print("Configuring Slide 3: Control Architecture...")
    for s in s3.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "Project data" in txt:
                s.text_frame.text = "Control Architecture: Two-Time-Scale Disturbance Observer"
                s.text_frame.paragraphs[0].font.size = Pt(19)
                s.text_frame.paragraphs[0].font.bold = True
            elif "Benefits of using tables" in txt:
                s.text_frame.text = "Mathematical Formulations"
                s.text_frame.paragraphs[0].font.size = Pt(13)
                s.text_frame.paragraphs[0].font.bold = True
            elif "Tables in project proposals" in txt:
                s.text_frame.clear()
                points = [
                    "• Observer Filter: dF_hat/dt = ω_obs · (F_meas - F_hat) with ω_obs = 35 rad/s.",
                    "• Feedforward Rejection: a_des = a_PID - (1/m) F_hat cancels wind drag.",
                    "• Tilt Thrust: T = m a_z / (cos φ cos θ) maintains rock-solid altitude during 37° lean.",
                    "• Lyapunov Stability: Error ||e_obs(t)|| → 0 exponentially under bounded gusts."
                ]
                for i, pt in enumerate(points):
                    p = s.text_frame.add_paragraph() if i > 0 else s.text_frame.paragraphs[0]
                    p.text = pt
                    p.font.size = Pt(9.5)
                    p.space_after = Pt(4)
            elif "Follow the link" in txt:
                s.text_frame.text = "Figure: Disturbance Observer Outer-Loop Feedforward & SO(3) Attitude Allocation"
                s.text_frame.paragraphs[0].font.size = Pt(8.5)

    # Replace picture on Slide 12 with architecture diagram
    replace_shape_image(s3, None, os.path.join(IMG_DIR, "control_system_architecture.png"))

    # -------------------------------------------------------------
    # 4. EDIT SLIDE 14 (CUSTOM_4 -> 3D GAZEBO SIMULATION ARENA)
    # -------------------------------------------------------------
    s4 = slides_list[13]
    print("Configuring Slide 4: 3D Gazebo Simulation Arena...")
    for s in s4.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "What to show in a mockup" in txt:
                s.text_frame.text = "3D Simulation Arena & 11-Waypoint Obstacle Corridor"
                s.text_frame.paragraphs[0].font.size = Pt(15)
                s.text_frame.paragraphs[0].font.bold = True
            elif "Product/website description" in txt:
                s.text_frame.clear()
                points = [
                    "• 6 Industrial Obstacles: Helipad, 4m Gate, 2m Rooftop Station, Towers A & B, Pylons.",
                    "• Verified Clearances: Margins ≥ 0.82 m - 2.10 m from rotor tip to any obstacle surface.",
                    "• Inspection Gate Fly-Through: Flies through opening (W=3.2m, H=4m) to Y=6.5m before turning.",
                    "• Rooftop Survey: Surveys elevated target at Z=3.5m (1.5m above roof) under 8.5 m/s crosswinds."
                ]
                for i, pt in enumerate(points):
                    p = s.text_frame.add_paragraph() if i > 0 else s.text_frame.paragraphs[0]
                    p.text = pt
                    p.font.size = Pt(9.5)
                    p.space_after = Pt(4)

    # Replace picture on Slide 14 with 3D Gazebo simulation arena
    replace_shape_image(s4, None, os.path.join(IMG_DIR, "gazebo_3d_simulation_arena.png"))

    # -------------------------------------------------------------
    # 5. EDIT SLIDE 18 (CUSTOM_22 -> MULTI-GUST EMPIRICAL BENCHMARK)
    # -------------------------------------------------------------
    s5 = slides_list[17]
    print("Configuring Slide 5: Multi-Gust Empirical Benchmark...")
    for s in s5.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "Project expenses" in txt:
                s.text_frame.text = "Empirical Results: Multi-Position Wind Gust Benchmark"
                s.text_frame.paragraphs[0].font.size = Pt(19)
                s.text_frame.paragraphs[0].font.bold = True
            elif "Labor" in txt:
                s.text_frame.text = "Gust 1 (Gate, 7.5m/s): 85.8% Cut (12.2cm vs 86.2cm)"
                s.text_frame.paragraphs[0].font.size = Pt(9.0)
            elif "Materials" in txt:
                s.text_frame.text = "Gust 2 (Rooftop, 8.5m/s): 73.3% Cut (41.9cm vs 157.1cm)"
                s.text_frame.paragraphs[0].font.size = Pt(9.0)
            elif "Equipment" in txt:
                s.text_frame.text = "Gust 3 (Helipad, 8.0m/s): 64.7% Cut (29.9cm vs 84.7cm)"
                s.text_frame.paragraphs[0].font.size = Pt(9.0)
            elif "Overhead" in txt:
                s.text_frame.text = "Global Cross-Track Error: 38.7% Reduction (20.2cm vs 32.9cm)"
                s.text_frame.paragraphs[0].font.size = Pt(9.0)
            elif "Other" in txt:
                s.text_frame.text = "Safety Envelope: Zero obstacle collisions across all 11 waypoints"
                s.text_frame.paragraphs[0].font.size = Pt(9.0)
            elif "Follow the link" in txt:
                s.text_frame.text = "Empirical Comparison: 3D Trajectory Tracking & Disturbance Reconstruction"
                s.text_frame.paragraphs[0].font.size = Pt(8.5)

    # Replace picture on Slide 18 with multi-gust benchmark plot
    replace_shape_image(s5, None, os.path.join(IMG_DIR, "multi_gust_dob_vs_pid.png"))

    # -------------------------------------------------------------
    # 6. EDIT SLIDE 19 (CUSTOM_22_1 -> GCS DASHBOARD & KPI)
    # -------------------------------------------------------------
    s6 = slides_list[18]
    print("Configuring Slide 6: Ground Control Station (GCS)...")
    for s in s6.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "KPI dashboard" in txt:
                s.text_frame.text = "Ground Control Station (GCS) & Human Teleoperation"
                s.text_frame.paragraphs[0].font.size = Pt(19)
                s.text_frame.paragraphs[0].font.bold = True
            elif "120 u/day" in txt:
                s.text_frame.text = "85.8%"
                s.text_frame.paragraphs[0].font.size = Pt(16)
                s.text_frame.paragraphs[0].font.bold = True
            elif "Output per worker" in txt:
                s.text_frame.text = "Peak Drift Reduction"
                s.text_frame.paragraphs[0].font.size = Pt(9.5)
            elif "2h" in txt:
                s.text_frame.text = "28.5 ms"
                s.text_frame.paragraphs[0].font.size = Pt(16)
                s.text_frame.paragraphs[0].font.bold = True
            elif "Time to complete" in txt:
                s.text_frame.text = "Observer Time Constant"
                s.text_frame.paragraphs[0].font.size = Pt(9.5)
            elif "Follow the link" in txt:
                s.text_frame.text = "Features: Dynamic Drag Calculator (Fd = ½ρCdAv²), 1-Click Patrol, Emergency Abort"
                s.text_frame.paragraphs[0].font.size = Pt(8.5)

    # Replace picture on Slide 19 with GCS Dashboard HUD
    replace_shape_image(s6, None, os.path.join(IMG_DIR, "gcs_teleop_dashboard.png"))

    # -------------------------------------------------------------
    # 7. EDIT SLIDE 20 (CUSTOM_3 -> THANKS & DEFENSE SUMMARY)
    # -------------------------------------------------------------
    s7 = slides_list[19]
    print("Configuring Slide 7: Thanks & Summary...")
    for s in s7.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "Thanks!" in txt:
                s.text_frame.text = "Summary & Q&A"
                s.text_frame.paragraphs[0].font.size = Pt(28)
                s.text_frame.paragraphs[0].font.bold = True
            elif "Does anyone have any questions?" in txt:
                s.text_frame.text = (
                    "• Core Contribution: Active Disturbance Rejection Control (ADRC) with proven Lyapunov stability.\n"
                    "• Verified Performance: 85.8% drift cut in Gazebo 11 physics under severe aerodynamic wind shocks.\n"
                    "• Complete Open-Source Suite: ROS 2 Humble packages, GCS HUD, and automated benchmark tests.\n\n"
                    "GitHub: github.com/Sarat-2007/ros2-quadrotor-dob-wind-rejection (Branch: main)"
                )
                p = s.text_frame.paragraphs[0]
                p.font.size = Pt(10.5)

    # -------------------------------------------------------------
    # REMOVE ALL OTHER 34 UNUSED BOILERPLATE / ICON SLIDES
    # -------------------------------------------------------------
    print("Pruning 34 unused boilerplate/template slides in reverse order...")
    for idx in range(total_initial - 1, -1, -1):
        if idx not in KEEP_INDICES:
            delete_slide(prs, idx)

    print(f"Final presentation slide count: {len(prs.slides)}")
    print(f"Saving presentation to: {OUT_PATH}")
    prs.save(OUT_PATH)
    print("Presentation saved successfully!")

if __name__ == "__main__":
    build()
