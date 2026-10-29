#!/usr/bin/env python3
from pptx import Presentation
prs = Presentation("/mnt/d/ROS/UAV_Disturbance_Observer_Project_Presentation.pptx")
print(f"Total Slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    texts = []
    has_image = False
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    texts.append(p.text.strip())
        if shape.shape_type == 13: # Picture
            has_image = True
    print(f"Slide {i+1}: {len(slide.shapes)} shapes | Has Real Image: {has_image} | Header: {texts[0] if texts else 'None'}")
