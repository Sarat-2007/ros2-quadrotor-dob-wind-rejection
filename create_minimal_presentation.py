#!/usr/bin/env python3
"""
create_minimal_presentation.py
Generates a minimal, clean, professional academic PowerPoint presentation (.pptx)
Key design principles:
- No crazy templates or garish colors (clean white background, deep navy & slate).
- Real simulation figures embedded with strict aspect-ratio fitting and centering (NO stretching).
- Calibrated layout geometry eliminating any overlapping text, cards, or images.
- Consistent typography and margins across all 8 slides.
"""

import os
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = "/mnt/d/ROS/UAV_Disturbance_Observer_Presentation.pptx"
IMG_DIR = "/mnt/d/ROS/docs/images"

# --- REFINED ACADEMIC COLOR PALETTE ---
COLOR_BG = RGBColor(255, 255, 255)         # Pure White
COLOR_NAVY = RGBColor(30, 41, 59)          # Deep Slate/Navy (#1E293B)
COLOR_STEEL = RGBColor(71, 85, 105)        # Muted Slate (#475569)
COLOR_TEXT = RGBColor(51, 65, 85)          # Charcoal (#334155)
COLOR_MUTED = RGBColor(148, 163, 184)      # Subdued Grey (#94A3B8)
COLOR_CARD_BG = RGBColor(248, 250, 252)    # Clean Slate Fill (#F8FAFC)
COLOR_CARD_BORDER = RGBColor(226, 232, 240)# Clean Slate Border (#E2E8F0)
COLOR_HIGHLIGHT = RGBColor(15, 118, 110)   # Deep Teal Accent (#0F766E)

def add_fitted_picture(slide, img_path, box_left, box_top, box_w, box_h):
    """Calculates aspect ratio, fits image inside bounding box without distortion, and centers it."""
    if not os.path.exists(img_path):
        return None
    
    with Image.open(img_path) as im:
        im_w, im_h = im.size
        im_ratio = im_w / im_h

    box_ratio = box_w / box_h

    if im_ratio > box_ratio:
        # Image is wider than box -> constrain by width
        fit_w = box_w
        fit_h = box_w / im_ratio
    else:
        # Image is taller than box -> constrain by height
        fit_h = box_h
        fit_w = box_h * im_ratio

    # Center inside the box
    offset_x = box_left + (box_w - fit_w) / 2.0
    offset_y = box_top + (box_h - fit_h) / 2.0

    # Add subtle border container around the image for clean finish
    border_rect = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(offset_x - 0.04), Inches(offset_y - 0.04), 
        Inches(fit_w + 0.08), Inches(fit_h + 0.08)
    )
    border_rect.fill.solid()
    border_rect.fill.fore_color.rgb = COLOR_CARD_BG
    border_rect.line.color.rgb = COLOR_CARD_BORDER
    border_rect.line.width = Pt(1.0)

    # Insert picture
    pic = slide.shapes.add_picture(
        img_path, 
        Inches(offset_x), Inches(offset_y), 
        Inches(fit_w), Inches(fit_h)
    )
    return pic

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_base_slide(title_text, category_tag="UAV AUTONOMOUS FLIGHT & ROBUST CONTROL"):
        slide = prs.slides.add_slide(blank_layout)
        
        # Subtle header category tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.25))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_tag.upper()
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(9.0)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_STEEL

        # Main slide title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.733), Inches(0.55))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY

        # Subtle divider rule
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.30), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_CARD_BORDER
        line.line.color.rgb = COLOR_CARD_BORDER

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.25))
        tf_foot = footer_box.text_frame
        tf_foot.word_wrap = True
        tf_foot.margin_left = tf_foot.margin_top = tf_foot.margin_right = tf_foot.margin_bottom = 0
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = "Autonomous Aerial Robotics Research Laboratory  |  ROS 2 Humble & Gazebo 11  |  github.com/Sarat-2007/ros2-quadrotor-dob-wind-rejection"
        p_foot.font.name = "Calibri"
        p_foot.font.size = Pt(8.5)
        p_foot.font.color.rgb = COLOR_MUTED

        return slide

    def add_card(slide, left, top, width, height, title, items, badge=""):
        # Card background
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        rect.fill.solid()
        rect.fill.fore_color.rgb = COLOR_CARD_BG
        rect.line.color.rgb = COLOR_CARD_BORDER
        rect.line.width = Pt(1.0)

        # Card text content
        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.22), Inches(width - 0.5), Inches(height - 0.44))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Arial"
        p0.font.size = Pt(12.5)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_NAVY
        p0.space_after = Pt(8)

        for item in items:
            p = tf.add_paragraph()
            p.space_after = Pt(5)
            p.font.name = "Calibri"
            p.font.size = Pt(10.5)
            p.font.color.rgb = COLOR_TEXT
            if isinstance(item, tuple):
                bold_prefix, text_body = item
                r_pre = p.add_run()
                r_pre.text = bold_prefix + " "
                r_pre.font.bold = True
                r_pre.font.color.rgb = COLOR_NAVY
                r_body = p.add_run()
                r_body.text = text_body
            else:
                r = p.add_run()
                r.text = item

    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE (Clean, Minimal, Modern)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    
    # Left accent indicator
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.0), Inches(0.06), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_NAVY
    bar.line.fill.background()

    tb_t = s1.shapes.add_textbox(Inches(1.25), Inches(1.9), Inches(11.0), Inches(3.4))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0

    p_proj = tf_t.paragraphs[0]
    p_proj.text = "RESEARCH & TECHNICAL EVALUATION PRESENTATION"
    p_proj.font.name = "Arial"
    p_proj.font.size = Pt(10)
    p_proj.font.bold = True
    p_proj.font.color.rgb = COLOR_STEEL
    p_proj.space_after = Pt(12)

    p_title = tf_t.add_paragraph()
    p_title.text = "Nonlinear Disturbance Observer (DOB) for Autonomous Quadrotors Under Severe Aerodynamic Wind Shocks"
    p_title.font.name = "Arial"
    p_title.font.size = Pt(23)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_NAVY
    p_title.space_after = Pt(10)

    p_sub = tf_t.add_paragraph()
    p_sub.text = "High-Fidelity ROS 2 Humble & Gazebo 11 Physics Simulation, Differential Flatness Allocation, and Empirical Benchmark"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = COLOR_STEEL
    p_sub.space_after = Pt(26)

    p_meta = tf_t.add_paragraph()
    p_meta.text = "Author: Nallabati Venkata Durga Sai Sarat Chandra  |  Department of Robotics & Autonomous Systems\nTarget Repository: github.com/Sarat-2007/ros2-quadrotor-dob-wind-rejection (Branch: main)"
    p_meta.font.name = "Calibri"
    p_meta.font.size = Pt(10.5)
    p_meta.font.color.rgb = COLOR_TEXT

    # -------------------------------------------------------------
    # SLIDE 2: THE PROBLEM - WHY CLASSICAL PID FAILS IN WIND
    # -------------------------------------------------------------
    s2 = add_base_slide("Problem Statement: Aerodynamic Wind Shocks vs. Classical PID", "Control Limitation Analysis")

    items_flaws = [
        ("The Inspection Challenge:", "UAVs conducting facility perimeter patrol maneuver within 1-2 meters of buildings, inspection gates, and transmission towers."),
        ("Rayleigh Fluid Drag:", "An 8.5 m/s shock crosswind exerts ~11.2 N on a 1.52 kg quadrotor—representing 75.1% of entire vehicle weight."),
        ("PID Error-Driven Lag:", "PID controllers only react after error has accumulated. Under shock winds, position drift exceeds 1.5 meters."),
        ("The Integral Windup Dilemma:", "Integral terms slowly wind up over seconds. Once wind ceases, accumulated energy causes massive overshoot into obstacles."),
        ("Actuator Saturation Risk:", "Increasing high PID gains amplifies sensor noise, causing motor chattering and rapid instability.")
    ]
    add_card(s2, 0.8, 1.5, 5.7, 5.2, "Classical PID Failure Modes", items_flaws)

    items_real = [
        ("Peak Crosswind Velocity:", "8.5 m/s Shock Step Injections in Gazebo 11 physics."),
        ("Baseline PID Drift:", "157.1 cm lateral displacement (Near-miss crash hazard)."),
        ("DOB Rejection Drift:", "41.9 cm maximum displacement (73.3% drift reduction)."),
        ("Observer Reaction Time:", "28.5 milliseconds (vs. 2.4 - 4.0 seconds for PID integrator)."),
        ("Altitude Loss Under Tilt:", "0 cm drop (Nonlinear tilt thrust compensation maintains Z)."),
        ("Collision Record:", "Zero obstacle collisions across all 11 facility waypoints.")
    ]
    add_card(s2, 6.8, 1.5, 5.7, 5.2, "Target Specifications & Comparative Metrics", items_real)

    # -------------------------------------------------------------
    # SLIDE 3: DISTURBANCE OBSERVER ARCHITECTURE & THEORY
    # -------------------------------------------------------------
    s3 = add_base_slide("Control Architecture: Two-Time-Scale Disturbance Observer", "Robust Control Theory")

    items_theory = [
        ("Observer State Filter:", "dF_hat/dt = ω_obs · (F_meas - F_hat), with bandwidth ω_obs = 35.0 rad/s."),
        ("Feedforward Decoupling:", "a_des = a_ref + K_p e_p + K_i ∫e_p + K_d e_v - (1/m) F_hat. Direct acceleration cancellation."),
        ("Nonlinear Tilt Thrust:", "T = (m · a_z,des) / (cos φ · cos θ). Prevents altitude drop when tilting up to 37° into wind."),
        ("Differential Flatness:", "Maps desired accelerations into SO(3) attitude angles (φ_des, θ_des) for the inner-loop rate controller."),
        ("Lyapunov Proof:", "Proves exponential convergence of disturbance estimation error: ||e_obs(t)|| → 0 as t → ∞.")
    ]
    add_card(s3, 0.8, 1.5, 5.3, 5.2, "Mathematical Formulations", items_theory)

    # Fit Real Architecture Image (16:9 ratio)
    add_fitted_picture(s3, os.path.join(IMG_DIR, "control_system_architecture.png"), 6.4, 1.5, 6.1, 5.2)

    # -------------------------------------------------------------
    # SLIDE 4: 3D GAZEBO ARENA & 11-WAYPOINT OBSTACLE CORRIDOR
    # -------------------------------------------------------------
    s4 = add_base_slide("3D Gazebo Simulation Arena & 11-Waypoint Flight Corridor", "Facility Environment Modeling")

    items_arena = [
        ("6 Realistic Obstacles:", "Central Helipad (Origin), 4m Inspection Gate, 2m Elevated Rooftop Station, Towers A & B, Slalom Pylons."),
        ("11-Waypoint Corridor:", "Designed with verified clearances (≥ 0.82 m - 2.10 m from prop tip to any physical surface)."),
        ("Inspection Gate Transit:", "Flies through center (W=3.2m, H=4m) to Y=6.5m before turning, eliminating pillar clipping."),
        ("Rooftop Station Survey:", "Surveys elevated target at Z=3.5m (1.5m above roof surface) under active crosswinds."),
        ("Multi-Zone Gust Points:", "Shock gusts injected at Gate Approach (7.5 m/s), Rooftop (8.5 m/s), and Helipad Recovery (8.0 m/s).")
    ]
    add_card(s4, 0.8, 1.5, 5.3, 5.2, "Safety Envelope & Waypoints", items_arena)

    # Fit Real 3D Arena Image (4:3 ratio)
    add_fitted_picture(s4, os.path.join(IMG_DIR, "gazebo_3d_simulation_arena.png"), 6.4, 1.5, 6.1, 5.2)

    # -------------------------------------------------------------
    # SLIDE 5: EMPIRICAL MULTI-GUST BENCHMARK RESULTS
    # -------------------------------------------------------------
    s5 = add_base_slide("Empirical Results: Multi-Position Wind Gust Benchmark", "Quantitative Performance Verification")

    items_res = [
        ("Gust 1 (Gate, 7.5 m/s):", "PID Drift = 86.2 cm  vs.  DOB Drift = 12.2 cm  (85.8% Reduction). Prevents gate pillar collision."),
        ("Gust 2 (Rooftop, 8.5 m/s):", "PID Drift = 157.1 cm  vs.  DOB Drift = 41.9 cm  (73.3% Reduction). Prevents vehicle roof blowout."),
        ("Gust 3 (Helipad, 8.0 m/s):", "PID Drift = 84.7 cm  vs.  DOB Drift = 29.9 cm  (64.7% Reduction). Guarantees 1.2m pad touchdown."),
        ("Overall Cross-Track Error:", "Reduced from 32.9 cm to 20.2 cm across the entire 11-waypoint patrol (38.7% overall improvement)."),
        ("Actuator Demands:", "Smooth motor thrust scaling (14.9 N → 18.2 N) without motor saturation or chattering.")
    ]
    add_card(s5, 0.8, 1.5, 5.3, 5.2, "Drift Reduction Highlights", items_res)

    # Fit Real Multi-Gust Benchmark Image
    add_fitted_picture(s5, os.path.join(IMG_DIR, "multi_gust_dob_vs_pid.png"), 6.4, 1.5, 6.1, 5.2)

    # -------------------------------------------------------------
    # SLIDE 6: OBSERVER FREQUENCY RESPONSE & MOTOR DYNAMICS
    # -------------------------------------------------------------
    s6 = add_base_slide("Disturbance Reconstruction & Actuator Dynamics", "Observer Convergence Analysis")

    items_dyn = [
        ("Observer Bandwidth (35 rad/s):", "Time constant τ = 28.5 ms enables near-instantaneous reconstruction of external aerodynamic drag."),
        ("Force Estimation Accuracy:", "Reconstructs 10.33 N shock force to 10.34 N (>98% fidelity) in time domain."),
        ("Separation Principle:", "Bandwidth sits safely between position loop (5 rad/s) and attitude loop (100 rad/s), preventing noise amplification."),
        ("Tilt Thrust Balancing:", "Collective thrust automatically increases during tilt, ensuring zero altitude drop during 35° leans."),
        ("Motor PWM Utilization:", "All 4 rotor speeds operate within 1200 - 1800 rad/s, maintaining safe margins from physical saturation.")
    ]
    add_card(s6, 0.8, 1.5, 5.3, 5.2, "Transient Dynamics & Bandwidth", items_dyn)

    # Fit Real Experiment Benchmark Image
    add_fitted_picture(s6, os.path.join(IMG_DIR, "experiment_benchmark_report.png"), 6.4, 1.5, 6.1, 5.2)

    # -------------------------------------------------------------
    # SLIDE 7: GROUND CONTROL STATION & PILOT INTERACTION
    # -------------------------------------------------------------
    s7 = add_base_slide("Ground Control Station (GCS) & Human-in-the-Loop", "Industrial Pilot Dashboard")

    items_gcs = [
        ("Industrial Dark Slate HUD:", "High-contrast telemetry interface showing live 3D coordinates, speed, heading, and disturbance wrench."),
        ("Aerodynamic Drag Calculator:", "Operator sets gust slider (1-12 m/s); UI computes expected drag (Fd), counter-tilt angle, and required thrust."),
        ("1-Click Patrol Execution:", "Deploys autonomous inspection sequence across all 11 waypoints with live progress tracking."),
        ("Emergency Station Hover:", "Instant mission abort halting trajectory tracking and locking drone into zero-drift station hover."),
        ("Pilot Manual Overrides:", "Touching any keyboard key or D-Pad stick automatically yields mission control to the human pilot.")
    ]
    # WIDER card on left (7.6 inches) to balance portrait image on right
    add_card(s7, 0.8, 1.5, 7.6, 5.2, "Operational Capabilities", items_gcs)

    # Portrait GCS Dashboard Image fitted on right
    add_fitted_picture(s7, os.path.join(IMG_DIR, "gcs_teleop_dashboard.png"), 8.7, 1.5, 3.8, 5.2)

    # -------------------------------------------------------------
    # SLIDE 8: SUMMARY & FACULTY EVALUATION TAKEAWAYS
    # -------------------------------------------------------------
    s8 = add_base_slide("Project Summary & Key Defense Takeaways", "Academic Evaluation Reference")

    items_sum_left = [
        ("Not Just a Waypoint Follower:", "Standard drone projects follow waypoints in calm air. This project introduces Active Disturbance Rejection Control (ADRC) for violent unmodeled wind conditions."),
        ("Rigorous Mathematical Basis:", "Formulated 6-DOF rigid-body dynamics, Rayleigh fluid drag, and proved Lyapunov asymptotic stability."),
        ("Exceptional Benchmark Gain:", "Demonstrated 85.8% drift reduction at Gate and 73.3% at Rooftop Station with zero collisions.")
    ]
    add_card(s8, 0.8, 1.5, 5.7, 5.2, "Core Technical Novelty", items_sum_left)

    items_sum_right = [
        ("Modular ROS 2 Architecture:", "Clean separation across uav_control (100 Hz), uav_aerodynamics, uav_dashboard, and uav_description."),
        ("Reproducible Open Source Suite:", "Automated 1-click test scripts and benchmarks enable instant verification for examiners."),
        ("Ready for Hardware Transition:", "Differential flatness and acceleration decoupling structure maps directly to PX4 / Pixhawk flight hardware.")
    ]
    add_card(s8, 6.8, 1.5, 5.7, 5.2, "Engineering Deliverables", items_sum_right)

    print(f"Saving presentation to {PPTX_PATH}...")
    prs.save(PPTX_PATH)
    print("Presentation saved successfully!")

if __name__ == "__main__":
    create_deck()
